"""PDF summaries and Excel lists. A PDF holds the few facts a planner needs; the Excel holds every row."""
import io
from datetime import datetime

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

INK, MUTED, ACCENT = colors.HexColor("#1f2937"), colors.HexColor("#6b7280"), colors.HexColor("#0f766e")


def _fmt(n, d=0):
    return f"{n:,.{d}f}"


# ------------------------------------------------------------------- PDF helpers
def _styles():
    s = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("t", parent=s["Title"], fontSize=18, alignment=0, textColor=INK, spaceAfter=2),
        "sub": ParagraphStyle("s", parent=s["Normal"], fontSize=9, textColor=MUTED, spaceAfter=10),
        "h": ParagraphStyle("h", parent=s["Heading2"], fontSize=12, textColor=INK, spaceBefore=12, spaceAfter=4),
        "p": ParagraphStyle("p", parent=s["Normal"], fontSize=9.5, leading=13, textColor=INK),
    }


def _table(data, widths=None, align_right_from=1):
    t = Table(data, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 8.5), ("TEXTCOLOR", (0, 0), (-1, -1), INK),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e6f2f0")), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#e5e7eb")),
        ("ALIGN", (align_right_from, 0), (-1, -1), "RIGHT"), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    return t


def _kpis(items):
    cells = [[Paragraph(f'<font size="15" color="#0f766e"><b>{v}</b></font><br/><font size="8" color="#6b7280">{k}</font>', getSampleStyleSheet()["Normal"]) for k, v in items]]
    t = Table(cells, colWidths=[(A4[0] - 40 * mm) / len(items)] * len(items))
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#d1d5db")),
                           ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#e5e7eb")), ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    return t


def _chart(fig) -> Image:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=160, bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return Image(buf, width=170 * mm, height=60 * mm, kind="proportional")


def _pdf(title, subtitle, story_fn) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=16 * mm, bottomMargin=16 * mm, title=title)
    st = _styles()
    story = [Paragraph(title, st["title"]), Paragraph(subtitle, st["sub"])]
    story += story_fn(st)
    doc.build(story)
    return buf.getvalue()


def _stamp(scope: str) -> str:
    return f"{scope} &nbsp;|&nbsp; Created {datetime.now().strftime('%d %b %Y, %H:%M')}"


def _line_chart(x, series: dict, ylabel="Units"):
    fig, ax = plt.subplots(figsize=(7.4, 2.6))
    for label, (ys, color, style) in series.items():
        ax.plot(x, ys, style, color=color, label=label, linewidth=1.8)
    ax.set_ylabel(ylabel, fontsize=8)
    ax.tick_params(labelsize=7)
    ax.grid(alpha=0.25)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    if len(series) > 1:
        ax.legend(fontsize=7, frameon=False)
    step = max(1, len(x) // 8)
    ax.set_xticks(range(0, len(x), step))
    ax.set_xticklabels([x[i][5:] for i in range(0, len(x), step)])
    return fig


# --------------------------------------------------------------- Excel helpers
def _sheet(wb, title, headers, rows, widths=None, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0F766E")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(11, min(34, len(h) + 4)))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def _xlsx(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------- order list
ORDER_HEAD = ["Priority", "Product ID", "Product", "Category", "Warehouse", "City", "Units to order", "Available now",
              "On order", "Days of cover", "Lead time (days)", "Cover period (days)", "Demand over cover period",
              "Demand (high end)", "On hand", "Reserved"]


def _order_row(x):
    return [x["priority"], x["product_id"], x["product_name"], x["category"], x["warehouse_id"], x["city"],
            x["units_to_order"], x["available"], x.get("on_order", 0), x["days_of_cover"], x.get("lead_time_days"),
            x.get("cover_days"), x["forecast_units"], x["forecast_high"], x["stock_on_hand"], x["reserved"]]


def orders_xlsx(res: dict) -> bytes:
    s = res["summary"]
    wb = Workbook()
    ws = _sheet(wb, "Summary", ["What", "Value"], [
        ["Planning period (days)", res["days"]], ["Default supplier lead time (days)", res.get("default_lead_time_days")],
        ["Warehouse", res["warehouse_id"] or "All warehouses"],
        ["Safety buffer", "On (upper end of the 80% forecast range)" if res["safety_buffer"] else "Off (expected demand)"],
        ["SKU-warehouse lines to order", s["skus_to_order"]], ["Total units to order", s["units_to_order"]],
        ["Order now", s["order_now"]], ["Order soon", s["order_soon"]],
        ["Units already on order", res.get("on_order_units", 0)], ["Sales data up to", res["history_to"]],
        ["Created", datetime.now().strftime("%Y-%m-%d %H:%M")],
    ], {"What": 38, "Value": 46}, first=True)
    ws.auto_filter.ref = None
    _sheet(wb, "Order list", ORDER_HEAD, [_order_row(x) for x in res["rows"]],
           {"Product": 34, "Category": 16, "Priority": 12})
    _sheet(wb, "By warehouse", ["Warehouse", "City", "SKUs to order", "Units to order"],
           [[w["warehouse_id"], w["city"], w["skus"], w["units"]] for w in res["by_warehouse"]])
    return _xlsx(wb)


def orders_pdf(res: dict) -> bytes:
    s = res["summary"]
    buffer_txt = "with a safety buffer" if res["safety_buffer"] else "without a safety buffer"

    def story(st):
        out = [_kpis([("Units to order", _fmt(s["units_to_order"])), ("SKU lines", _fmt(s["skus_to_order"])),
                      ("Order now", _fmt(s["order_now"])), ("Order soon", _fmt(s["order_soon"]))]), Spacer(1, 6)]
        out.append(Paragraph(
            f"Units needed so the shelf lasts until the next delivery: demand over the supplier's lead time plus the next "
            f"<b>{res['days']} days</b> {buffer_txt}, less stock free to sell and less the "
            f"{_fmt(res.get('on_order_units', 0))} units already on order. Lines that are already covered are not listed.", st["p"]))
        if res["by_warehouse"]:
            out.append(Paragraph("By warehouse", st["h"]))
            out.append(_table([["Warehouse", "City", "SKUs", "Units"]] + [[w["warehouse_id"], w["city"], w["skus"], _fmt(w["units"])] for w in res["by_warehouse"]],
                              [30 * mm, 50 * mm, 30 * mm, 40 * mm], 2))
        top = res["rows"][:25]
        out.append(Paragraph(f"Most urgent {len(top)} lines" if top else "Nothing to order", st["h"]))
        if top:
            out.append(_table([["Priority", "Product", "WH", "Order", "Have", "On order", "Lead (d)", "Cover (d)"]] +
                              [[x["priority"], (x["product_name"] or "")[:30], x["warehouse_id"], _fmt(x["units_to_order"]),
                                _fmt(x["available"]), _fmt(x.get("on_order", 0)), x.get("lead_time_days", "-"),
                                "-" if x["days_of_cover"] is None else x["days_of_cover"]] for x in top],
                              [22 * mm, 52 * mm, 14 * mm, 18 * mm, 18 * mm, 20 * mm, 17 * mm, 19 * mm], 3))
            if len(res["rows"]) > len(top):
                out.append(Paragraph(f"The Excel list has all {len(res['rows'])} lines.", st["sub"]))
        return out

    scope = res["warehouse_id"] or "All warehouses"
    return _pdf(f"Units to order - next {res['days']} days",
                _stamp(f"{scope} &nbsp;|&nbsp; lead time included &nbsp;|&nbsp; sales data up to {res['history_to']}"), story)


# ----------------------------------------------------------------- forecast
def forecast_xlsx(res: dict) -> bytes:
    wb = Workbook()
    s = res["summary"]
    ws = _sheet(wb, "Summary", ["What", "Value"], [
        ["Forecast period (days)", res["days"]], ["Warehouse", res["warehouse_id"] or "All warehouses"],
        ["Total forecast units", s["forecast_units"]], ["SKU-warehouse lines", s["skus_checked"]], ["Sales data up to", res["history_to"]],
    ], {"What": 30, "Value": 30}, first=True)
    ws.auto_filter.ref = None
    _sheet(wb, "By SKU", ORDER_HEAD[1:6] + ["Forecast demand", "Forecast (high end)", "On hand", "Available now", "Days of cover"],
           [[x["product_id"], x["product_name"], x["category"], x["warehouse_id"], x["city"], x["forecast_units"], x["forecast_high"],
             x["stock_on_hand"], x["available"], x["days_of_cover"]] for x in sorted(res["rows"], key=lambda r: -r["forecast_units"])],
           {"Product": 34})
    _sheet(wb, "By day", ["Date", "Forecast units"], [[d["date"], d["units"]] for d in res["daily"]])
    return _xlsx(wb)


def forecast_pdf(res: dict) -> bytes:
    s = res["summary"]
    rows = res["rows"]
    wh: dict = {}
    for x in rows:
        w = wh.setdefault(x["warehouse_id"], {"city": x["city"], "units": 0.0})
        w["units"] += x["forecast_units"]
    risky = [x for x in rows if x["days_of_cover"] is not None and x["days_of_cover"] < res["days"]]

    def story(st):
        out = [_kpis([("Forecast demand", _fmt(s["forecast_units"])), ("SKU lines", _fmt(s["skus_checked"])),
                      ("Will run short", _fmt(len(risky))), ("Days", str(res["days"]))]), Spacer(1, 6)]
        days = [d["date"] for d in res["daily"]]
        out.append(Paragraph("Expected demand per day", st["h"]))
        out.append(_chart(_line_chart(days, {"Forecast": ([d["units"] for d in res["daily"]], "#0f766e", "-")})))
        if len(wh) > 1:
            out.append(Paragraph("By warehouse", st["h"]))
            out.append(_table([["Warehouse", "City", "Forecast units"]] + [[k, v["city"], _fmt(v["units"])] for k, v in sorted(wh.items(), key=lambda kv: -kv[1]["units"])],
                              [35 * mm, 60 * mm, 45 * mm], 2))
        top = sorted(rows, key=lambda r: -r["forecast_units"])[:10]
        out.append(Paragraph("Highest demand SKUs", st["h"]))
        out.append(_table([["Product", "WH", "Forecast", "Have", "Cover (d)"]] + [[(x["product_name"] or "")[:36], x["warehouse_id"], _fmt(x["forecast_units"]),
                          _fmt(x["available"]), "-" if x["days_of_cover"] is None else x["days_of_cover"]] for x in top], [70 * mm, 18 * mm, 24 * mm, 24 * mm, 24 * mm], 2))
        out.append(Paragraph("Stock that will run out within the period is counted in &quot;Will run short&quot;. The order list report says how many units to buy.", st["sub"]))
        return out

    scope = res["warehouse_id"] or "All warehouses"
    return _pdf(f"Demand forecast - next {res['days']} days", _stamp(f"{scope} &nbsp;|&nbsp; sales data up to {res['history_to']}"), story)


# ----------------------------------------------------------------- accuracy
def _verdict(k: dict, target: float) -> str:
    m = k["day_mape_pct"]
    bias = k["day_bias_pct"]
    lean = "slightly under-forecasts" if bias < -0.5 else "slightly over-forecasts" if bias > 0.5 else "is balanced"
    ok = "within" if m <= target else "above"
    return (f"Daily forecasts are off by about {m:.0f}% on average, {ok} the {target:.0f}% target. Over a week per SKU the miss is about "
            f"{k['sku_week_mape_pct']:.0f}%. On balance the model {lean} ({bias:+.1f}%).")


def accuracy_xlsx(rep: dict) -> bytes:
    k = rep["kpis"]
    wb = Workbook()
    ws = _sheet(wb, "Summary", ["Measure", "Value"], [
        ["Scope", rep.get("warehouse_id") or "All warehouses"], ["Period", f"{rep['period']['start']} to {rep['period']['end']}"],
        ["Daily error (MAPE %)", k["day_mape_pct"]], ["Daily error (WAPE)", k["day_wape"]], ["Weekly SKU error (MAPE %)", k["sku_week_mape_pct"]],
        ["Bias %", k["day_bias_pct"]], ["Target MAPE %", rep["target_mape_pct"]], ["Total predicted", k["total_predicted"]], ["Total actual", k["total_actual"]],
        ["Reading", _verdict(k, rep["target_mape_pct"])],
    ], {"Measure": 30, "Value": 70}, first=True)
    ws.auto_filter.ref = None
    _sheet(wb, "Daily", ["Date", "Predicted", "Actual"], [[d["date"], d["predicted"], d["actual"]] for d in rep["daily"]])
    _sheet(wb, "By forecast day", list(rep["by_horizon"][0].keys()) if rep["by_horizon"] else ["day"], [list(r.values()) for r in rep["by_horizon"]])
    _sheet(wb, "By warehouse", list(rep["by_warehouse"][0].keys()), [list(r.values()) for r in rep["by_warehouse"]])
    _sheet(wb, "By SKU", list(rep["by_sku"][0].keys()), [list(r.values()) for r in rep["by_sku"]])
    return _xlsx(wb)


def accuracy_pdf(rep: dict) -> bytes:
    k = rep["kpis"]

    def story(st):
        out = [_kpis([("Daily error", f"{k['day_mape_pct']:.1f}%"), ("Weekly SKU error", f"{k['sku_week_mape_pct']:.1f}%"),
                      ("Bias", f"{k['day_bias_pct']:+.1f}%"), ("Target", f"{rep['target_mape_pct']:.0f}%")]), Spacer(1, 6),
               Paragraph(_verdict(k, rep["target_mape_pct"]), st["p"]), Paragraph("Predicted vs actual", st["h"])]
        d = rep["daily"]
        out.append(_chart(_line_chart([r["date"] for r in d], {"Actual": ([r["actual"] for r in d], "#374151", "-"),
                                                              "Predicted": ([r["predicted"] for r in d], "#0f766e", "--")})))
        if len(rep["by_warehouse"]) > 1:
            out.append(Paragraph("By warehouse", st["h"]))
            out.append(_table([["Warehouse", "Error %", "Bias %"]] + [[w["warehouse_id"], f"{w['mape_pct']:.1f}", f"{w['bias_pct']:+.1f}"] for w in rep["by_warehouse"]],
                              [40 * mm, 35 * mm, 35 * mm], 1))
        worst = sorted([s for s in rep["by_sku"] if s["mean_abs_err_pct"] is not None], key=lambda s: -s["mean_abs_err_pct"])[:8]
        out.append(Paragraph("SKUs with the largest weekly miss", st["h"]))
        out.append(_table([["Product", "WH", "Forecast", "Actual", "Miss %"]] + [[str(s.get("product_name") or s["product_id"])[:36], s["warehouse_id"],
                          _fmt(s["forecast_7d"]), _fmt(s["actual_7d"]), f"{s['mean_abs_err_pct']:.0f}"] for s in worst], [70 * mm, 18 * mm, 24 * mm, 24 * mm, 24 * mm], 2))
        return out

    scope = rep.get("warehouse_id") or "All warehouses"
    return _pdf("Forecast accuracy", _stamp(f"{scope} &nbsp;|&nbsp; {rep['period']['start']} to {rep['period']['end']}"), story)
